"""Prepare the user-supplied, unfilled reference as a runtime template bundle."""
import sys, json, base64, zipfile, re
from pathlib import Path
from docx import Document
source, output = map(Path, sys.argv[1:3])
doc = Document(source)
blocks=[]; pi=ti=0
for el in doc.element.body:
    kind=el.tag.rsplit('}',1)[-1]
    if kind=='p':
        blocks.append({'kind':'p','index':pi,'text':doc.paragraphs[pi].text,'xml':el.xml}); pi+=1
    elif kind=='tbl':
        blocks.append({'kind':'table','index':ti,'rows':[[c.text for c in row.cells] for row in doc.tables[ti].rows],'xml':el.xml}); ti+=1
    elif kind=='sectPr': section=el.xml
with zipfile.ZipFile(source) as archive:
    parts={n:base64.b64encode(archive.read(n)).decode() for n in archive.namelist() if n not in ('word/document.xml','docProps/core.xml')}
    document=archive.read('word/document.xml').decode('utf-8')
parts['docProps/core.xml']=base64.b64encode(b'<?xml version="1.0" encoding="UTF-8"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"/>').decode()
output.parent.mkdir(parents=True,exist_ok=True)
output.write_text(json.dumps({'version':'space-contract-20261025-v1','root':document.split('<w:body>')[0]+'<w:body>','section':section,'blocks':blocks,'parts':parts},ensure_ascii=False,separators=(',',':')),encoding='utf-8')
print('Prepared unfilled template: %d paragraphs, %d tables'%(pi,ti))
